"""Build the QPS house templates from a CODEX Word master.

Usage: python3 tools/house_style/build_templates.py <Word master .docx> <authoritative deck .pptx> [outdir]
  e.g. Full_VCR_Handover_master.docx ~/QPS_MTBF_WCS_DMAIC_v10.pptx   (outdir defaults to templates/)

Outputs
  QPS_House_starter.docx  Word reference doc (styles, numbering, corporate header/footer, empty body)
  QPS_House.dotx          same, as a Word template
  QPS_House.pptx          SCK CEN 2020 slide master (all layouts, no slides)
  QPS_House.potx          same, as a PowerPoint template
"""
import os
import sys
import zipfile
from pathlib import Path

import docx
from pptx import Presentation

HERE = Path(__file__).resolve().parent
TEMPLATES = HERE / "templates"
sys.path.insert(0, str(HERE))


def swap_content_type(src, dst, old, new):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(old, new)
            zout.writestr(it, data)


def build_word(master, out):
    d = docx.Document(master)
    body = d.element.body
    for el in list(body):
        if not el.tag.endswith("sectPr"):
            body.remove(el)
    starter = os.path.join(out, "QPS_House_starter.docx")
    d.save(starter)
    swap_content_type(starter, os.path.join(out, "QPS_House.dotx"),
                      b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")


def build_slides(deck, out):
    """Corporate slide template = the SCK CEN 2020 master from an authoritative deck, slides removed."""
    p = Presentation(deck)
    ids = p.slides._sldIdLst
    for sid in list(ids):
        p.part.drop_rel(sid.rId)
        ids.remove(sid)
    # Name the master furniture so the post-processor can fill it.
    m = p.slide_masters[0]
    boxes = sorted((s for s in m.shapes if s.shape_type == 17), key=lambda s: (s.left, s.top))
    for shp in boxes:
        txt = shp.text_frame.text.strip()
        if txt.startswith("ISC"):
            shp.name = "QPS ISC"
        elif txt.startswith("SCK CEN/"):
            shp.name = "QPS Ref"
        else:
            shp.name = "QPS Footer"
    # Google Slides exports substitute fonts (e.g. Quattrocento Sans); point them back at the theme font.
    from qps_post import normalise_fonts
    for part in list(p.slide_masters) + [lay for mm in p.slide_masters for lay in mm.slide_layouts]:
        normalise_fonts(part.element)
    pptx_path = os.path.join(out, "QPS_House.pptx")
    p.save(pptx_path)
    swap_content_type(pptx_path, os.path.join(out, "QPS_House.potx"),
                      b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")


if __name__ == "__main__":
    master, deck = sys.argv[1], sys.argv[2]
    outdir = sys.argv[3] if len(sys.argv) > 3 else str(TEMPLATES)
    os.makedirs(outdir, exist_ok=True)
    build_word(master, outdir)
    build_slides(deck, outdir)
