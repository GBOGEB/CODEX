"""QPS post-processing layer: always-on corporate finish for documents and decks.

  docx: Markdown/HTML/ODT/DOCX -> pandoc (--reference-doc QPS_House_starter.docx)
        -> MINERVA front matter (title block, copyright, configuration control,
        table of contents, abbreviations) in its own section -> corporate
        header/footer fields filled (name, reference, classification, date).
  pptx: Markdown -> pandoc, or a PPTX from any tool -> re-hosted onto the
        SCK CEN 2020 master from QPS_MTBF_WCS_DMAIC_v10 (Text Slide / BLANK
        layouts, logo, reference, ISC, slide numbers) -> cover on "Titel Slide 1".

Usage
  python3 qps_post.py docx IN OUT --title T [--subtitle S] [--project P] [--author A]
                      [--ref R] [--third-party-ref X] [--version V] [--status S]
                      [--classification C] [--date YYYY-MM-DD] [--short] [--keep] [--pdf]
  python3 qps_post.py pptx IN OUT --title T [same options] [--pdf]

--short  title block only, no configuration-control pages (informal notes)
--keep   (docx) skip the pandoc pass for a .docx already built on the starter
--pdf    also export a PDF beside OUT
"""
import argparse
import copy
import datetime as dt
import io
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
import docx.shared
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

TEMPLATES = Path(__file__).resolve().parent / "templates"
DOCX_REF = str(TEMPLATES / "QPS_House_starter.docx")
PPTX_REF = str(TEMPLATES / "QPS_House.pptx")  # built locally by build_templates.py; not committed
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PNS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

CORPORATE_SCHEME = "SCK•CEN 2020"
STANDARD_FONTS = {"Arial", "Segoe UI", "Segoe UI Semibold", "Segoe UI Light", "Symbol", "Wingdings", "Cambria Math"}
COVER, CONTENT, BLANK = "Titel Slide 1", "Text Slide", "BLANK"
# Content area of the "Text Slide" layout body and the full usable area above the footer (EMU).
CONTENT_BOX = (766763, 1486463, 10658474, 4860000)
BLANK_BOX = (766763, 760000, 10658474, 5590000)

COPYRIGHT = ("All property rights and copyright are reserved to SCK CEN. In case of a contractual "
             "arrangement with SCK CEN the use of this information by a Third Party or for any purpose "
             "other than for which it is intended on the basis of the contract is not authorized.")
IMPRINT = [
    "© SCK CEN - Stichting van Openbaar Nut – Fondation d'Utilité Publique - Foundation of Public Utility",
    "Registered Office: Avenue Herrmann Debroux 40 – BE‐1160 BRUSSEL",
    "Research Centres:",
    "Boeretang 200 - 2400 MOL – Belgium",
    "Chemin du Cyclotron 6 - 1348 Ottignies-Louvain-la-Neuve – Belgium",
    "www.sckcen.be",
]


def pandoc(src, dst, ref=None):
    args = ["pandoc", src, "-o", dst]
    if ref:
        args += ["--reference-doc", ref]
    subprocess.run(args, check=True)


# ----------------------------------------------------------------- docx ----

def _style(d, name, fallback="Normal"):
    try:
        return d.styles[name]
    except KeyError:
        return d.styles[fallback]


def _grid(d, header, rows):
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t.style = d.styles["Table Grid"]
    for i, h in enumerate(header):
        t.cell(0, i).text = h
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    house_table(t)
    return t


def _toc(d):
    p = d.add_paragraph()
    def fld(kind):
        r = etree.SubElement(p._p, f"{{{W}}}r")
        etree.SubElement(r, f"{{{W}}}fldChar").set(f"{{{W}}}fldCharType", kind)
    fld("begin")
    r = etree.SubElement(p._p, f"{{{W}}}r")
    it = etree.SubElement(r, f"{{{W}}}instrText"); it.text = ' TOC \\o "1-3" \\h \\z \\u '
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    fld("separate")
    p.add_run("Update this field (F9) to build the table of contents.")
    fld("end")
    return p


def _front_matter(d, meta):
    """MINERVA front matter, built with the house styles."""
    body = d.element.body
    for p in list(d.paragraphs)[:6]:
        if p.style.name in ("Title", "Subtitle", "Author", "Date", "Abstract"):
            body.remove(p._p)
    first = body[0]
    made = [d.add_paragraph(f"SCK CEN Ref.: {meta.ref or 'SCK CEN/########'}"),
            d.add_paragraph(f"Third Party Ref.: {meta.third_party_ref or '-'}"),
            d.add_paragraph(meta.title, style="Title")]
    if meta.subtitle:
        made.append(d.add_paragraph(meta.subtitle, style="Subtitle"))
    if meta.project:
        made.append(d.add_paragraph(meta.project))
    made.append(d.add_paragraph(" · ".join(x for x in (meta.author, meta.date,
                                 meta.version and f"Revision {meta.version}", f"ISC: {meta.classification}") if x)))
    if not meta.short:
        made.append(_section_break(d))
        hd = _style(d, "Heading without Listing", "Heading 1")
        made.append(d.add_paragraph("Copyright Rules", style=hd))
        made.append(d.add_paragraph(COPYRIGHT))
        made.append(d.add_paragraph("Configuration Control", style=hd))
        made.append(_subhead(d, "Document History"))
        made.append(_grid(d, ["Revision", "Status", "Author", "Date"],
                          [[meta.version or "0.1", meta.status, meta.author, meta.date], ["", "", "", ""]]))
        made.append(_subhead(d, "Document Approval"))
        made.append(_grid(d, ["Name", "Date", "Signature"], [["", "", ""], ["", "", ""]]))
        made.append(_subhead(d, "Distribution List"))
        made.append(_grid(d, ["Name", "Expert Group/Unit", "Number"], [["", "", ""], ["", "", ""]]))
        for line in IMPRINT:
            made.append(d.add_paragraph(line, style=_style(d, "No Spacing")))
        made.append(_section_break(d))
        made.append(d.add_paragraph("Table of Contents", style=hd))
        made.append(_toc(d))
        made.append(d.add_paragraph("Abbreviations, glossary and symbols", style=hd))
        made.append(_grid(d, ["Abbreviation", "Meaning"], [["", ""]]))
    made.append(_section_break(d))
    for el in made:
        first.addprevious(getattr(el, "_tbl", None) if hasattr(el, "_tbl") else getattr(el, "_p", el))


def _subhead(d, text):
    """Front-matter sub-heading: H2 look without joining the chapter numbering."""
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = docx.shared.RGBColor.from_string("562873")
    return p


def _section_break(d):
    """An empty paragraph closing a section (next page) with the body's header/footer."""
    p = d.add_paragraph()._p
    sect = copy.deepcopy(d.element.body.find(f"{{{W}}}sectPr"))
    for old in sect.findall(f"{{{W}}}type"):
        sect.remove(old)
    kind = etree.Element(f"{{{W}}}type"); kind.set(f"{{{W}}}val", "nextPage")
    sect.insert(0, kind)
    p.get_or_add_pPr().append(sect)
    return p


def house_table(t):
    """Full-width grid with a purple header row (white bold text)."""
    tblpr = t._tbl.tblPr
    for old in tblpr.findall(f"{{{W}}}tblW") + tblpr.findall(f"{{{W}}}tblLayout"):
        tblpr.remove(old)
    # Keep the schema order: tblStyle, tblW, jc, ..., tblLayout, tblCellMar, tblLook.
    tw = etree.Element(f"{{{W}}}tblW"); tw.set(f"{{{W}}}type", "pct"); tw.set(f"{{{W}}}w", "5000")
    style = tblpr.find(f"{{{W}}}tblStyle")
    if style is not None:
        style.addnext(tw)
    else:
        tblpr.insert(0, tw)
    layout = etree.Element(f"{{{W}}}tblLayout"); layout.set(f"{{{W}}}type", "autofit")
    anchor = tblpr.find(f"{{{W}}}tblCellMar")
    if anchor is None:
        anchor = tblpr.find(f"{{{W}}}tblLook")
    if anchor is not None:
        anchor.addprevious(layout)
    else:
        tblpr.append(layout)
    for jc in tblpr.findall(f"{{{W}}}jc"):
        if jc.get(f"{{{W}}}val") == "start":  # strict-OOXML value pandoc writes; Word 2007/LibreOffice want "left"
            jc.set(f"{{{W}}}val", "left")
    for jc in t._tbl.iter(f"{{{W}}}jc"):
        if jc.get(f"{{{W}}}val") == "start":
            jc.set(f"{{{W}}}val", "left")
    for tcw in list(t._tbl.iter(f"{{{W}}}tcW")):
        tcw.getparent().remove(tcw)
    grid = t._tbl.find(f"{{{W}}}tblGrid")
    cols = grid.findall(f"{{{W}}}gridCol") if grid is not None else []
    for gc in cols:  # equal share of the A4 text width (~9,500 twips)
        gc.set(f"{{{W}}}w", str(9500 // max(len(cols), 1)))
    # Cell paragraphs: Normal with no spacing or indent. The master's document defaults carry a
    # hanging indent that "No Spacing"/"Compact" inherit, which squeezes cell text to one word per line.
    for para in t._tbl.iter(f"{{{W}}}p"):
        ppr = para.find(f"{{{W}}}pPr")
        if ppr is None:
            ppr = etree.Element(f"{{{W}}}pPr"); para.insert(0, ppr)
        st = ppr.find(f"{{{W}}}pStyle")
        if st is None or st.get(f"{{{W}}}val") in ("Compact", "NoSpacing", "BodyText", "FirstParagraph"):
            if st is None:
                st = etree.Element(f"{{{W}}}pStyle"); ppr.insert(0, st)
            st.set(f"{{{W}}}val", "Normal")
        for tag in ("spacing", "ind"):
            for old in ppr.findall(f"{{{W}}}{tag}"):
                ppr.remove(old)
        sp = etree.Element(f"{{{W}}}spacing"); sp.set(f"{{{W}}}before", "40"); sp.set(f"{{{W}}}after", "40")
        ind = etree.Element(f"{{{W}}}ind"); ind.set(f"{{{W}}}left", "0"); ind.set(f"{{{W}}}firstLine", "0")
        anchor = ppr.find(f"{{{W}}}jc")
        for el in (sp, ind):
            if anchor is not None:
                anchor.addprevious(el)
            else:
                ppr.append(el)
    for cell in t.rows[0].cells:
        tcpr = cell._tc.get_or_add_tcPr()
        if tcpr.find(f"{{{W}}}shd") is None:
            shd = etree.SubElement(tcpr, f"{{{W}}}shd")
            shd.set(f"{{{W}}}val", "clear"); shd.set(f"{{{W}}}color", "auto"); shd.set(f"{{{W}}}fill", "562873")
        for r in cell._tc.iter(f"{{{W}}}r"):
            rpr = r.find(f"{{{W}}}rPr")
            if rpr is None:
                rpr = etree.Element(f"{{{W}}}rPr"); r.insert(0, rpr)
            if rpr.find(f"{{{W}}}color") is not None:
                continue
            etree.SubElement(rpr, f"{{{W}}}b")
            etree.SubElement(rpr, f"{{{W}}}color").set(f"{{{W}}}val", "FFFFFF")


def _fill_fields(part_xml, values):
    """Set cached results of complex fields (DOCPROPERTY/SAVEDATE) in a header/footer."""
    root = etree.fromstring(part_xml)
    stack = []  # [instr, state, written]
    for el in root.iter(f"{{{W}}}fldChar", f"{{{W}}}instrText", f"{{{W}}}t"):
        tag = etree.QName(el).localname
        if tag == "fldChar":
            kind = el.get(f"{{{W}}}fldCharType")
            if kind == "begin":
                stack.append(["", "instr", False])
            elif kind == "separate" and stack:
                stack[-1][1] = "result"
            elif kind == "end" and stack:
                stack.pop()
        elif tag == "instrText" and stack:
            stack[-1][0] += el.text or ""
        elif tag == "t" and stack and stack[-1][1] == "result":
            instr = " ".join(stack[-1][0].split())
            if instr.startswith("IF"):
                continue
            for key, val in values.items():
                if key in instr and val is not None:
                    el.text = "" if stack[-1][2] else val
                    stack[-1][2] = True
                    break
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _custom_props(values):
    ns = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
    vt = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    root = etree.Element(f"{{{ns}}}Properties", nsmap={None: ns, "vt": vt})
    for i, (k, v) in enumerate(values.items(), start=2):
        prop = etree.SubElement(root, f"{{{ns}}}property",
                                fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}", pid=str(i), name=k)
        etree.SubElement(prop, f"{{{vt}}}lpwstr").text = v
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def post_docx(src, out, meta, keep=False):
    with tempfile.TemporaryDirectory() as tmp:
        work = src
        if not (keep and src.endswith(".docx")):
            work = os.path.join(tmp, "p.docx")
            pandoc(src, work, DOCX_REF)
        d = docx.Document(work)
        for t in d.tables:
            t.style = d.styles["Table Grid"]
            house_table(t)
        _front_matter(d, meta)
        d.core_properties.title = meta.title
        d.core_properties.author = meta.author or ""
        staged = os.path.join(tmp, "s.docx")
        d.save(staged)

        fields = {"DOCPROPERTY Name": meta.title,
                  "Short Reference": meta.ref or "",
                  "Information Security Classification": meta.classification,
                  "SAVEDATE": meta.date}
        props = {"Name": meta.title,
                 "Common Attributes_Short Reference": meta.ref or "",
                 "Common Attributes_Information Security Classification": meta.classification}
        # Close the staged archive before the temp dir is removed: Windows
        # cannot delete a file that is still open (WinError 32).
        with zipfile.ZipFile(staged) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            names = zin.namelist()
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename.startswith(("word/header", "word/footer")) and it.filename.endswith(".xml"):
                    data = _fill_fields(data, fields)
                elif it.filename == "docProps/custom.xml":
                    data = _custom_props(props)
                elif it.filename == "word/settings.xml" and b"updateFields" not in data:
                    # Ask Word to refresh the TOC and fields on open.
                    data = data.replace(b"</w:settings>", b'<w:updateFields w:val="true"/></w:settings>')
                zout.writestr(it, data)
            if "docProps/custom.xml" not in names:
                zout.writestr("docProps/custom.xml", _custom_props(props))


# ----------------------------------------------------------------- pptx ----

def is_corporate(prs):
    theme = prs.slide_masters[0].part.part_related_by(RT.THEME).blob
    return f'clrScheme name="{CORPORATE_SCHEME}"'.encode() in theme


def _layout(prs, name):
    for m in prs.slide_masters:
        for l in m.slide_layouts:
            if l.name == name:
                return l
    raise KeyError(name)


def _copy_rels(src_slide, dst_slide, el):
    for node in el.iter():
        for attr in (f"{{{R}}}embed", f"{{{R}}}link", f"{{{R}}}id"):
            rid = node.get(attr)
            if not rid or rid not in src_slide.part.rels:
                continue
            rel = src_slide.part.rels[rid]
            if rel.is_external:
                new = dst_slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif rel.reltype == RT.IMAGE:
                _, new = dst_slide.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            else:
                new = dst_slide.part.relate_to(rel.target_part, rel.reltype)
            node.set(attr, new)


def _explicit_xfrm(shp, el):
    """Give a copied placeholder its effective geometry and make it a plain shape."""
    for ph in el.xpath(".//p:nvPr/p:ph"):
        ph.getparent().remove(ph)
    sppr = el.find(f"{{{PNS}}}spPr")
    if sppr is not None and sppr.find(f"{{{A}}}xfrm") is None and shp.left is not None:
        xfrm = etree.Element(f"{{{A}}}xfrm")
        etree.SubElement(xfrm, f"{{{A}}}off", x=str(shp.left), y=str(shp.top))
        etree.SubElement(xfrm, f"{{{A}}}ext", cx=str(shp.width), cy=str(shp.height))
        sppr.insert(0, xfrm)


def _fit(elements, box):
    """Uniformly scale and move top-level shapes so their bounding box fits `box`."""
    offs = []
    for el in elements:
        xfrm = el.find(f"{{{PNS}}}spPr/{{{A}}}xfrm")
        if xfrm is None:
            xfrm = el.find(f"{{{PNS}}}xfrm")  # graphicFrame
        if xfrm is None:
            xfrm = el.find(f"{{{PNS}}}grpSpPr/{{{A}}}xfrm")
        if xfrm is None:
            continue
        off, ext = xfrm.find(f"{{{A}}}off"), xfrm.find(f"{{{A}}}ext")
        if off is not None and ext is not None:
            offs.append((off, ext))
    if not offs:
        return
    x0 = min(int(o.get("x")) for o, _ in offs); y0 = min(int(o.get("y")) for o, _ in offs)
    x1 = max(int(o.get("x")) + int(e.get("cx")) for o, e in offs)
    y1 = max(int(o.get("y")) + int(e.get("cy")) for o, e in offs)
    bx, by, bw, bh = box
    s = min(bw / max(x1 - x0, 1), bh / max(y1 - y0, 1), 1.0)  # shrink to fit, never enlarge
    for o, e in offs:
        o.set("x", str(int(bx + (int(o.get("x")) - x0) * s)))
        o.set("y", str(int(by + (int(o.get("y")) - y0) * s)))
        e.set("cx", str(int(int(e.get("cx")) * s)))
        e.set("cy", str(int(int(e.get("cy")) * s)))


def _ensure_slide_number(slide):
    if any(s.is_placeholder and s.placeholder_format.type == 13 for s in slide.shapes):
        return
    for s in slide.slide_layout.placeholders:
        if s.placeholder_format.type == 13:
            slide.shapes._spTree.append(copy.deepcopy(s._element))
            return


def _drop_empty_placeholders(slide):
    for s in list(slide.placeholders):
        if s.placeholder_format.type == 13:
            continue
        if s.has_text_frame and not s.text_frame.text.strip():
            s._element.getparent().remove(s._element)


def _title_of(slide):
    t = slide.shapes.title
    return t if t is not None and t.has_text_frame and t.text_frame.text.strip() else None


def rehost(src_prs, dst):
    for sl in src_prs.slides:
        title = _title_of(sl)
        layout = _layout(dst, CONTENT if title is not None else BLANK)
        new = dst.slides.add_slide(layout)
        copied = []
        if title is not None:
            new.shapes.title.text_frame.text = title.text_frame.text
        bodies = [s for s in sl.placeholders if s.placeholder_format.type in (2, 7) and s.has_text_frame
                  and s.text_frame.text.strip()]
        target = next((p for p in new.placeholders if p.placeholder_format.idx == 2), None)
        use_body = title is not None and len(bodies) == 1 and target is not None
        if use_body:
            # Move the text into the layout's body so it takes the corporate text styles.
            txb = copy.deepcopy(bodies[0]._element.find(f"{{{PNS}}}txBody"))
            old = target._element.find(f"{{{PNS}}}txBody")
            target._element.replace(old, txb)
            target.left, target.top, target.width, target.height = (
                target.left, target.top, target.width, target.height)  # pin geometry on the slide
        for shp in sl.shapes:
            if (title is not None and shp.shape_id == title.shape_id) or (use_body and shp.shape_id == bodies[0].shape_id):
                continue
            if shp.is_placeholder and shp.placeholder_format.type in (13, 15, 16):  # number, footer, date
                continue
            el = copy.deepcopy(shp._element)
            if shp.is_placeholder:
                _explicit_xfrm(shp, el)
            _copy_rels(sl, new, el)
            new.shapes._spTree.append(el)
            copied.append(el)
        box = CONTENT_BOX if title is not None else BLANK_BOX
        if use_body and copied:
            box = (box[0] + box[2] // 2 + 150000, box[1] + 800000, box[2] // 2 - 150000, box[3] - 800000)
            target.width = box[2]
        if use_body:
            shrink_to_fit(target)
        _fit(copied, box)
        _drop_empty_placeholders(new)
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame is not None:
            new.notes_slide.notes_text_frame.text = sl.notes_slide.notes_text_frame.text


def text_fit_scale(shape, default_pt=16):
    """Rough fraction of the text that fits in its shape (1.0 = fits)."""
    if not shape.has_text_frame or not shape.width or not shape.height:
        return 1.0
    width_pt, height_pt = shape.width / 12700, shape.height / 12700
    used = 0.0
    for p in shape.text_frame.paragraphs:
        size = next((r.font.size.pt for r in p.runs if r.font.size), default_pt)
        chars = max(len(p.text), 1)
        per_line = max(width_pt / (0.5 * size), 1)
        used += -(-chars // per_line) * size * 1.25
    return min(1.0, height_pt / used) if used else 1.0


def shrink_to_fit(shape):
    """Set shrink-on-overflow with a font scale PowerPoint and LibreOffice both honour."""
    scale = text_fit_scale(shape)
    body = shape.text_frame._txBody.find(f"{{{A}}}bodyPr")
    for old in body.findall(f"{{{A}}}normAutofit") + body.findall(f"{{{A}}}spAutoFit") + body.findall(f"{{{A}}}noAutofit"):
        body.remove(old)
    fit = etree.SubElement(body, f"{{{A}}}normAutofit")
    if scale < 1.0:
        pct = max(int(scale * 100000) // 2500 * 2500, 40000)
        fit.set("fontScale", str(pct))
        fit.set("lnSpcReduction", "10000")


def _set_text(shape, text):
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        shape.text_frame.text = text


def normalise_fonts(root):
    """Point off-standard Latin fonts (e.g. Google's Quattrocento Sans substitution) back at the theme font."""
    for el in root.iter(f"{{{A}}}latin"):
        tf = el.get("typeface", "")
        if tf and not tf.startswith("+") and tf not in STANDARD_FONTS:
            el.set("typeface", "+mn-lt")


def finish_deck(dst, meta):
    """Cover slide, slide numbers and master furniture on a corporate deck."""
    first = dst.slides[0] if len(dst.slides) else None
    if first is None or not first.slide_layout.name.startswith("Titel Slide"):
        cover = dst.slides.add_slide(_layout(dst, COVER))
        for ph in cover.placeholders:
            if ph.placeholder_format.type == 1:
                ph.text_frame.text = meta.title
            elif ph.placeholder_format.idx == 1:
                ph.text_frame.text = " · ".join(x for x in (meta.subtitle, meta.project, meta.date) if x) or meta.date
        ids = dst.slides._sldIdLst
        ids.insert(0, ids[-1])
    for sl in list(dst.slides)[1:]:
        _ensure_slide_number(sl)
    for sl in dst.slides:
        normalise_fonts(sl.shapes._spTree)
    m = dst.slide_masters[0]
    for shp in m.shapes:
        if shp.name == "QPS Ref" and meta.ref:
            _set_text(shp, meta.ref)
        elif shp.name == "QPS ISC":
            _set_text(shp, f"ISC: {meta.classification}")
        elif shp.name == "QPS Footer":
            _set_text(shp, meta.title)


def post_pptx(src, out, meta):
    with tempfile.TemporaryDirectory() as tmp:
        if not src.endswith(".pptx"):
            work = os.path.join(tmp, "p.pptx")
            pandoc(src, work)
            src = work
        src_prs = Presentation(src)
        if is_corporate(src_prs):
            dst = src_prs
        else:
            dst = Presentation(PPTX_REF)
            rehost(src_prs, dst)
        finish_deck(dst, meta)
        dst.save(out)


def add_meta_args(ap):
    ap.add_argument("--title", required=True)
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--project", default="")
    ap.add_argument("--author", default="Gerkotze Bonthuys")
    ap.add_argument("--ref", default="")
    ap.add_argument("--third-party-ref", default="")
    ap.add_argument("--version", default="")
    ap.add_argument("--status", default="Draft")
    ap.add_argument("--classification", default="Restricted")
    ap.add_argument("--date", default=dt.datetime.now(tz=dt.UTC).date().isoformat())
    ap.add_argument("--short", action="store_true")


def export_pdf(path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                    os.path.dirname(os.path.abspath(path)), path], check=True, capture_output=True)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("kind", choices=["docx", "pptx"])
    ap.add_argument("src")
    ap.add_argument("out")
    add_meta_args(ap)
    ap.add_argument("--keep", action="store_true")
    ap.add_argument("--pdf", action="store_true")
    a = ap.parse_args(argv)
    if a.kind == "docx":
        post_docx(a.src, a.out, a, keep=a.keep)
    else:
        post_pptx(a.src, a.out, a)
    if a.pdf:
        export_pdf(a.out)


if __name__ == "__main__":
    sys.exit(main())
