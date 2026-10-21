"""QPS publish-to-corporate QA: check a Word file or deck against the house standard, report, fix.

The standard is set by the authoritative references:
  slides  QPS_MTBF_WCS_DMAIC_v7.pptx (SCK CEN 2020 master, layouts, palette) -> templates/QPS_House.pptx
  Word    Full_VCR_Handover_master.docx (styles, header/footer) + MINERVA_Template_Enhanced.docx (front matter)

Usage
  python3 qps_qa.py check FILE [--json REPORT.json] [--md REPORT.md]
  python3 qps_qa.py fix FILE OUT --title T [qps_post options] [--pdf [--pdf-engine E]]   # fix, then re-check OUT

Exit code of `check`: 0 = publishable (no FAIL), 1 = at least one FAIL.
"""
import argparse
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
from lxml import etree
from pptx import Presentation

try:
    from . import qps_post
except ImportError:  # run as a script
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import qps_post

A, W = qps_post.A, qps_post.W

PALETTE = {  # SCK CEN 2020 theme
    "562873": "accent1 purple", "984A9C": "accent2 pink", "8ED8F8": "accent3 cyan",
    "034694": "accent4 navy", "CACCD0": "accent5 grey", "DFE0E2": "accent6 light grey",
    "515151": "dk2 text", "E7E6E6": "lt2", "BC58AE": "hyperlink", "000000": "black", "FFFFFF": "white",
}
FONTS = qps_post.STANDARD_FONTS
CORPORATE_LAYOUTS = None  # filled from QPS_House.pptx
WORD_STYLES = None        # filled from QPS_House_starter.docx
FRONT_MATTER = ["Copyright Rules", "Configuration Control", "Document History", "Document Approval",
                "Distribution List", "Table of Contents"]
MANUAL_NUMBER = re.compile(r"^\s*(§+\s*\d|\d+(\.\d+)*\.?\s+\S|[IVX]+\.\s)")
# Template text that is never legitimate content.
PLACEHOLDER_TEXT = re.compile(r"\[(ref|type of document|cluster|project|sections)\]|#{4,}|lorem ipsum", re.IGNORECASE)
# Author markers. Engineering content quotes them as data (W363 burndown `status: TODO`, "TBD mm2"), so they
# count only outside quotes and status fields, and on slides only in placeholders (title, body).
AUTHOR_MARKER = re.compile(r"\b(TODO|TBD|XXX)\b")
QUOTED_OR_STATUS = re.compile(r"\"[^\"\n]*\"|“[^”\n]*”|‘[^’\n]*’|«[^»\n]*»|`[^`\n]*`|(?<!\w)'[^'\n]*'(?!\w)"
                              r"|\bstatus\s*[:=]\s*\w+", re.IGNORECASE)


def placeholder_text(text, author_markers=True):
    """The template text or author marker left in `text`, or None."""
    m = PLACEHOLDER_TEXT.search(text)
    if m:
        return m.group(0)
    if author_markers:
        m = AUTHOR_MARKER.search(QUOTED_OR_STATUS.sub(" ", text))
        if m:
            return m.group(0)
    return None


def _neutral(hexv):
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    return max(r, g, b) - min(r, g, b) <= 12


def nearest(hexv):
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    def dist(p):
        pr, pg, pb = (int(p[i:i + 2], 16) for i in (0, 2, 4))
        return math.sqrt((r - pr) ** 2 + (g - pg) ** 2 + (b - pb) ** 2)
    return min((p for p in PALETTE if p not in ("000000", "FFFFFF")), key=dist)


def _tint_of_palette(hexv, tol=10):
    """True when the colour is a lighter tint (mix with white) of a palette colour."""
    c = [int(hexv[i:i + 2], 16) for i in (0, 2, 4)]
    for p in PALETTE:
        pc = [int(p[i:i + 2], 16) for i in (0, 2, 4)]
        ts = [(c[k] - pc[k]) / (255 - pc[k]) for k in range(3) if pc[k] < 250]
        if not ts:
            continue
        t = sum(ts) / len(ts)
        if 0 <= t <= 1 and all(abs(pc[k] + t * (255 - pc[k]) - c[k]) <= tol for k in range(3)):
            return True
    return False


def colour_ok(hexv):
    hexv = hexv.upper()
    return hexv in PALETTE or _neutral(hexv) or _tint_of_palette(hexv)


class Report:
    def __init__(self, path):
        self.path, self.items = path, []

    def add(self, level, rule, where, detail, fixable=True):
        self.items.append({"level": level, "rule": rule, "where": where, "detail": detail, "fixable": fixable})

    @property
    def ok(self):
        return not any(i["level"] == "FAIL" for i in self.items)

    def summary(self):
        c = {k: sum(1 for i in self.items if i["level"] == k) for k in ("FAIL", "WARN")}
        return f"{os.path.basename(self.path)}: {'PUBLISHABLE' if self.ok else 'NOT PUBLISHABLE'} ({c['FAIL']} fail, {c['WARN']} warn)"

    def markdown(self):
        lines = [f"# QA report: {os.path.basename(self.path)}", "", f"**{self.summary()}**", ""]
        if self.items:
            lines += ["| Level | Rule | Where | Detail | Auto-fix |", "|---|---|---|---|---|"]
            for i in self.items:
                lines.append(f"| {i['level']} | {i['rule']} | {i['where']} | {i['detail']} | {'yes' if i['fixable'] else 'no'} |")
        return "\n".join(lines) + "\n"


def _load_standard():
    global CORPORATE_LAYOUTS, WORD_STYLES
    if WORD_STYLES is None:
        WORD_STYLES = {s.name for s in docx.Document(qps_post.DOCX_REF).styles}
    if CORPORATE_LAYOUTS is None:
        if not Path(qps_post.PPTX_REF).exists():
            raise FileNotFoundError(
                f"{qps_post.PPTX_REF} is missing: build it with build_templates.py from the authoritative deck")
        CORPORATE_LAYOUTS = {l.name for m in Presentation(qps_post.PPTX_REF).slide_masters for l in m.slide_layouts}


# ----------------------------------------------------------------- pptx ----

def custom_props(prs):
    """docProps/custom.xml as {name: text}."""
    part = qps_post._docprops(prs, "custom.xml")
    if part is None:
        return {}
    return {p.get("name"): "".join(p.itertext()).strip() for p in etree.fromstring(part.blob)}


def _check_identity(prs, rep):
    """One footer reference on every master, and no ECM identity of another document in the properties."""
    refs = {}
    for partname, role, shp in qps_post.furniture(prs):
        if role == "ref":
            refs.setdefault(shp.text_frame.text.strip(), []).append(str(partname))
    if refs and not any(re.search(r"\d", r) for r in refs):
        rep.add("WARN", "reference", "master", "document reference in the footer is empty", False)
    if len(refs) > 1:
        rep.add("FAIL", "reference", "masters", "masters show different references: " + "; ".join(
            f"'{r}' on {len(parts)} part(s)" for r, parts in sorted(refs.items())))
    ecm_id = custom_props(prs).get("ID", "")
    if ecm_id and refs and not all(re.search(rf"(?<!\d){re.escape(ecm_id)}$", r) for r in refs):
        rep.add("FAIL", "identity", "document properties",
                f"ECM document ID {ecm_id} in docProps/custom.xml belongs to another document "
                f"(footer reference: {', '.join(sorted(refs))})")


def check_pptx(path):
    _load_standard()  # needs the locally built slide master
    rep = Report(path)
    prs = Presentation(path)
    if not qps_post.is_corporate(prs):
        rep.add("FAIL", "theme", "deck", f"theme is not '{qps_post.CORPORATE_SCHEME}'")
    if (prs.slide_width, prs.slide_height) != (12192000, 6858000):
        rep.add("FAIL", "format", "deck", "slide size is not 16:9 widescreen (33.867 × 19.05 cm)")
    slides = list(prs.slides)
    if not slides:
        rep.add("FAIL", "content", "deck", "deck has no slides", False)
        return rep
    if not slides[0].slide_layout.name.startswith("Titel Slide"):
        rep.add("FAIL", "title slide", "slide 1", "first slide is not a corporate title slide (Titel Slide 1-6)")
    _check_identity(prs, rep)
    for n, sl in enumerate(slides, start=1):
        where = f"slide {n}"
        if sl.slide_layout.name not in CORPORATE_LAYOUTS:
            rep.add("FAIL", "layout", where, f"layout '{sl.slide_layout.name}' is not a corporate layout")
        if n > 1:
            if not any(s.is_placeholder and s.placeholder_format.type == 13 for s in sl.shapes):
                rep.add("FAIL", "pagination", where, "no slide number")
            if qps_post._title_of(sl) is None and not qps_post.is_blank_layout(sl.slide_layout.name):
                rep.add("WARN", "content", where, "slide has no title")
        fonts, colours = set(), set()
        for el in sl.shapes._spTree.iter(f"{{{A}}}latin"):
            tf = el.get("typeface", "")
            if tf and not tf.startswith("+") and tf not in FONTS:
                fonts.add(tf)
        for el in sl.shapes._spTree.iter(f"{{{A}}}srgbClr"):
            if not colour_ok(el.get("val")):
                colours.add(el.get("val").upper())
        if fonts:
            rep.add("FAIL", "font", where, "off-standard fonts: " + ", ".join(sorted(fonts)))
        if colours:
            rep.add("WARN", "colour", where, "off-palette colours: " + ", ".join(
                f"#{c}→#{nearest(c)}" for c in sorted(colours)))
        for shp in sl.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                found = placeholder_text(shp.text_frame.text, author_markers=shp.is_placeholder)
                if found:
                    rep.add("FAIL", "content", where, f"template placeholder text '{found}' left in '{shp.name}'", False)
                if qps_post.text_fit_scale(shp) < 0.7:
                    body = shp.text_frame._txBody.find(f"{{{A}}}bodyPr")
                    fit = body.find(f"{{{A}}}normAutofit") if body is not None else None
                    if fit is None or not fit.get("fontScale"):
                        rep.add("WARN", "overflow", where, f"text in '{shp.name}' likely overflows its box")
    return rep


def fix_pptx(path, out, meta):
    _load_standard()
    prs = Presentation(path)
    if not qps_post.is_corporate(prs) or any(
            sl.slide_layout.name not in CORPORATE_LAYOUTS for sl in prs.slides):
        dst = Presentation(qps_post.PPTX_REF)
        qps_post.rehost(prs, dst)
    else:
        dst = prs
    qps_post.finish_deck(dst, meta)
    for sl in dst.slides:
        tree = sl.shapes._spTree
        qps_post.normalise_fonts(tree)
        for el in tree.iter(f"{{{A}}}srgbClr"):
            if not colour_ok(el.get("val")):
                el.set("val", nearest(el.get("val")))
        for shp in sl.shapes:
            if shp.has_text_frame and qps_post.text_fit_scale(shp) < 0.8:
                qps_post.shrink_to_fit(shp)
    dst.save(out)


# ----------------------------------------------------------------- docx ----

def _footer_has_fields(path):
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if re.match(r"word/footer\d*\.xml", n)]
        text = b"".join(z.read(n) for n in parts)
    return b"DOCPROPERTY" in text and b"Information Security Classification" in text, text


def _load_word_standard():
    global WORD_STYLES
    if WORD_STYLES is None:
        WORD_STYLES = {s.name for s in docx.Document(qps_post.DOCX_REF).styles}


def check_docx(path):
    _load_word_standard()
    rep = Report(path)
    d = docx.Document(path)
    has_fields, footer = _footer_has_fields(path)
    if not has_fields:
        rep.add("FAIL", "header/footer", "document", "corporate footer (name, reference, ISC fields) missing")
    elif b"Addendum II" in footer:
        rep.add("FAIL", "header/footer", "footer", "footer still shows the master's document name")
    sec = d.sections[-1]
    if abs(sec.page_width - 7560310) > 20000 or abs(sec.page_height - 10692130) > 20000:
        rep.add("FAIL", "format", "page", "page is not A4 portrait")
    paras = d.paragraphs
    texts = [p.text.strip() for p in paras]
    if not any(p.style.name == "Title" and p.text.strip() for p in paras[:15]):
        rep.add("FAIL", "title page", "start", "no title on the first page")
    if not any(t.startswith("SCK CEN Ref.") for t in texts[:10]):
        rep.add("WARN", "reference", "title page", "no 'SCK CEN Ref.' line on the title page")
    missing = [h for h in FRONT_MATTER if h not in texts]
    if missing:
        rep.add("FAIL", "front matter", "MINERVA", "missing: " + ", ".join(missing))
    if not any(p.style.name.startswith("Heading") for p in paras):
        rep.add("WARN", "content", "body", "no headings: document has no structure")
    off_styles, fonts, colours, manual = set(), set(), set(), []
    for i, p in enumerate(paras):
        if p.style.name not in WORD_STYLES:
            off_styles.add(p.style.name)
        if p.style.name.startswith("Heading") and MANUAL_NUMBER.match(p.text):
            manual.append(p.text[:40])
        if placeholder_text(p.text):
            rep.add("FAIL", "content", f"paragraph {i + 1}", f"template placeholder text: '{p.text[:50]}'", False)
    body = d.element.body
    for rf in body.iter(f"{{{W}}}rFonts"):
        f = rf.get(f"{{{W}}}ascii")
        if f and f not in FONTS:
            fonts.add(f)
    for c in body.iter(f"{{{W}}}color"):
        v = c.get(f"{{{W}}}val")
        if v and v != "auto" and not colour_ok(v):
            colours.add(v.upper())
    if off_styles:
        rep.add("WARN", "style", "body", "styles not in the master: " + ", ".join(sorted(off_styles)))
    if fonts:
        rep.add("FAIL", "font", "body", "direct off-standard fonts: " + ", ".join(sorted(fonts)))
    if colours:
        rep.add("WARN", "colour", "body", "off-palette colours: " + ", ".join(f"#{c}→#{nearest(c)}" for c in sorted(colours)))
    if manual:
        rep.add("WARN", "numbering", "headings", "typed numbers in auto-numbered headings: " + "; ".join(manual[:5]))
    return rep


def fix_docx(path, out, meta):
    rep = check_docx(path)
    rules = {i["rule"] for i in rep.items}
    with tempfile.TemporaryDirectory() as tmp:
        src = path
        if rules & {"header/footer", "format", "style"} or "front matter" in rules or "title page" in rules:
            # Re-flow onto the house reference doc and add the corporate front matter.
            keep = not (rules & {"header/footer", "format", "style"})
            staged = os.path.join(tmp, "post.docx")
            qps_post.post_docx(src, staged, meta, keep=keep)
            src = staged
        d = docx.Document(src)
        body = d.element.body
        for rf in list(body.iter(f"{{{W}}}rFonts")):
            if rf.get(f"{{{W}}}ascii") and rf.get(f"{{{W}}}ascii") not in FONTS:
                rf.getparent().remove(rf)
        for c in body.iter(f"{{{W}}}color"):
            v = c.get(f"{{{W}}}val")
            if v and v != "auto" and not colour_ok(v):
                c.set(f"{{{W}}}val", nearest(v))
        for p in d.paragraphs:
            if p.style.name.startswith("Heading") and MANUAL_NUMBER.match(p.text) and p.runs:
                p.runs[0].text = MANUAL_NUMBER.sub(lambda m: m.group(0)[-1] if m.group(0)[-1].isalpha() else "", p.runs[0].text, 1).lstrip()
        tmp_out = os.path.join(tmp, "fixed.docx")
        d.save(tmp_out)
        shutil.copy(tmp_out, out)


# ----------------------------------------------------------------- cli -----

def check(path):
    return check_pptx(path) if path.lower().endswith(".pptx") else check_docx(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check"); c.add_argument("file"); c.add_argument("--json"); c.add_argument("--md")
    f = sub.add_parser("fix"); f.add_argument("file"); f.add_argument("out"); qps_post.add_meta_args(f)
    qps_post.add_pdf_args(f)
    a = ap.parse_args(argv)
    if a.cmd == "check":
        rep = check(a.file)
    else:
        (fix_pptx if a.file.lower().endswith(".pptx") else fix_docx)(a.file, a.out, a)
        if a.pdf:
            qps_post.export_pdf(a.out, a.pdf_engine)
        rep = check(a.out)
        a.json = a.md = None
    print(rep.markdown())
    if getattr(a, "json", None):
        Path(a.json).write_text(json.dumps({"file": a.file, "publishable": rep.ok, "items": rep.items}, indent=2))
    if getattr(a, "md", None):
        Path(a.md).write_text(rep.markdown())
    return 0 if rep.ok else 1


if __name__ == "__main__":
    sys.exit(main())
