"""Coverage wave 15 (L3): QPS round-trip semantic manifest extractor and
comparator (07_ops/qps_roundtrip).

``defusedxml`` is not installed in CI, so the extractor is loaded with a shim
that maps ``defusedxml.ElementTree`` to the stdlib parser (same ``fromstring``
API); fixtures are generated with openpyxl / python-docx / python-pptx.
"""
from __future__ import annotations

import json
import sys
import types
import xml.etree.ElementTree as StdET

import pytest

from tests.coverage_waves._loader import ROOT, run_main

GET = "07_ops/qps_roundtrip/Get-QpsSemanticManifest.py"
CMP = "07_ops/qps_roundtrip/Compare-QpsSemanticManifests.py"


@pytest.fixture
def get_mod(monkeypatch):
    import importlib.util
    shim = types.ModuleType("defusedxml")
    shim.ElementTree = StdET
    monkeypatch.setitem(sys.modules, "defusedxml", shim)
    monkeypatch.setitem(sys.modules, "defusedxml.ElementTree", StdET)
    spec = importlib.util.spec_from_file_location("l3cw_get_qps_semantic_manifest", ROOT / GET)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture
def cmp_mod():
    import importlib.util
    spec = importlib.util.spec_from_file_location("l3cw_compare_qps_semantic_manifests", ROOT / CMP)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------- helpers
def test_norm_text_and_hash(get_mod):
    m = get_mod
    assert m.norm_text("  a \t b\r\nc\r ") == "a b\nc"
    assert m.norm_text(None) == ""
    assert m.sha256_text(" x ") == m.sha256_text("x")
    assert m.resolve_part("xl/workbook.xml", "worksheets/sheet1.xml") == "xl/worksheets/sheet1.xml"
    assert m.resolve_part("xl/workbook.xml", "/xl/sheet.xml") == "xl/sheet.xml"


def test_xlsx_cell_value_variants(get_mod):
    m = get_mod
    ns = m.NS["s"]

    def cell(xml):
        return StdET.fromstring(f'<c xmlns="{ns}" {xml}')

    shared = ["zero", "one"]
    assert m.xlsx_cell_value(cell('t="inlineStr"><is><t> hi </t></is></c>'), shared) == "hi"
    assert m.xlsx_cell_value(cell('t="inlineStr"></c>'), shared) == ""
    assert m.xlsx_cell_value(cell("></c>"), shared) is None
    assert m.xlsx_cell_value(cell('t="s"><v>1</v></c>'), shared) == "one"
    assert m.xlsx_cell_value(cell('t="s"><v>9</v></c>'), shared) == "9"
    assert m.xlsx_cell_value(cell('t="b"><v>1</v></c>'), shared) == "TRUE"
    assert m.xlsx_cell_value(cell('t="b"><v>0</v></c>'), shared) == "FALSE"
    assert m.xlsx_cell_value(cell('t="b"><v>x</v></c>'), shared) == "x"
    assert m.xlsx_cell_value(cell("><v>3.5</v></c>"), shared) == "3.5"


# ---------------------------------------------------------------- per-format manifests
def test_xlsx_manifest(get_mod, tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "label"
    ws["B1"] = 2
    ws["C1"] = "=B1*2"
    ws["D1"] = True
    wb.create_sheet("Second")["A1"] = "x"
    from openpyxl.workbook.defined_name import DefinedName
    try:
        wb.defined_names["RATE"] = DefinedName("RATE", attr_text="Data!$B$1")
    except TypeError:
        wb.defined_names.append(DefinedName("RATE", attr_text="Data!$B$1"))
    path = tmp_path / "m.xlsx"
    wb.save(path)
    res = get_mod.build(path)
    assert res["kind"] == "xlsx" and res["sheet_order"] == ["Data", "Second"]
    assert res["normalized_formula_map"]["Data!C1"] == "B1*2"
    vals = res["normalized_value_map_for_controlled_cells"]
    assert vals["Data!A1"] == "label" and vals["Data!B1"] == "2" and vals["Data!D1"] == "TRUE"
    assert {"name": "RATE", "value": "Data!$B$1"} in res["defined_names"]
    assert res["artifact_name"] == "m.xlsx" and len(res["semantic_sha256"]) == 64


def test_docx_manifest(get_mod, tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_heading("Scope", level=1)
    d.add_paragraph("Body text")
    d.add_paragraph("")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "a"
    path = tmp_path / "d.docx"
    d.save(path)
    res = get_mod.build(path)
    assert res["kind"] == "docx" and res["heading_sequence"] == ["Scope"]
    assert len(res["paragraph_text_hashes"]) >= 2
    assert res["table_shape_and_text_hashes"][0]["rows"] == 2 and res["table_shape_and_text_hashes"][0]["cols_max"] == 2


def test_pptx_manifest_with_notes(get_mod, tmp_path):
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Overview"
    s.notes_slide.notes_text_frame.text = "speaker note"
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "p.pptx"
    prs.save(path)
    res = get_mod.build(path)
    assert res["kind"] == "pptx" and res["slide_titles"] == ["Overview", ""]
    assert len(res["slide_order"]) == 2 and res["object_count_by_slide"][0] >= 1
    assert len(res["notes_text_hashes"]) == 2


@pytest.mark.xfail(
    strict=True,
    reason=(
        "production_defect: Get-QpsSemanticManifest.resolve_part joins relative targets with PurePosixPath, "
        "which does not collapse '..', so PowerPoint's '../notesSlides/notesSlideN.xml' never resolves and "
        "every speaker-notes hash is the empty-string hash"
    ),
)
def test_pptx_manifest_captures_speaker_notes(get_mod, tmp_path):
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.notes_slide.notes_text_frame.text = "speaker note"
    path = tmp_path / "n.pptx"
    prs.save(path)
    assert get_mod.build(path)["notes_text_hashes"][0] == get_mod.sha256_text("speaker note")


def test_html_manifest(get_mod, tmp_path):
    path = tmp_path / "x.html"
    path.write_text('<div id="sec1"><h1>Title <b>X</b></h1><img src="a.png"><a href="b.html">b</a>'
                    '<script type="application/json">{"k":1}</script><style>p{}</style><p>Visible</p></div>',
                    encoding="utf-8")
    res = get_mod.build(path)
    assert res["route_or_section_ids"] == ["sec1"] and res["heading_sequence"] == ["Title X"]
    assert res["asset_references"] == ["a.png", "b.html"] and len(res["embedded_data_hashes"]) == 1


def test_build_rejects_unsupported(get_mod, tmp_path):
    with pytest.raises(SystemExit, match="Unsupported artifact type: .txt"):
        get_mod.build(tmp_path / "a.txt")


def test_xml_helpers_missing_parts(get_mod, tmp_path):
    import zipfile
    z = tmp_path / "empty.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("x.txt", "x")
    with zipfile.ZipFile(z) as zf:
        assert get_mod.xml_root(zf, "nope.xml") is None
        assert get_mod.relmap(zf, "nope.rels") == {} and get_mod.relationship_targets(zf, "nope.rels") == []
        assert get_mod.xlsx_shared_strings(zf) == [] and get_mod.pptx_notes_path(zf, "ppt/slides/s1.xml") is None
    assert get_mod.pptx_object_count(None) == 0
    assert get_mod.pptx_object_count(StdET.fromstring("<x/>")) == 0
    shell = tmp_path / "shell.xlsx"
    with zipfile.ZipFile(shell, "w") as zf:
        zf.writestr("x.txt", "x")
    assert get_mod.xlsx_manifest(shell)["sheet_order"] == []
    assert get_mod.docx_manifest(shell)["heading_sequence"] == []
    assert get_mod.pptx_manifest(shell)["slide_order"] == []


def test_get_main_stdout_and_file(get_mod, tmp_path, monkeypatch, capsys):
    html = tmp_path / "x.html"
    html.write_text("<h2>A</h2>", encoding="utf-8")
    monkeypatch.setattr("sys.argv", ["g", str(html)])
    get_mod.main()
    assert json.loads(capsys.readouterr().out)["heading_sequence"] == ["A"]
    out = tmp_path / "o" / "m.json"
    monkeypatch.setattr("sys.argv", ["g", str(html), "-o", str(out)])
    get_mod.main()
    assert json.loads(out.read_text(encoding="utf-8"))["kind"] == "html"


# ---------------------------------------------------------------- comparator
def test_canonicalize_and_hash(cmp_mod):
    m = cmp_mod
    a = m.canonicalize({"b": [{"generated_utc": 1, "x": 2}], "a": 1, "artifact_sha256": "z"})
    assert a == {"a": 1, "b": [{"x": 2}]}
    assert m.sha256_json({"a": 1, "b": 2}) == m.sha256_json({"b": 2, "a": 1})


def test_diff_reasons(cmp_mod):
    m = cmp_mod
    d = m.diff({"a": 1, "b": [1, 2], "c": {"x": 1}, "t": 1}, {"a": 2, "b": [1], "d": 1, "c": {"x": 1}, "t": "1"})
    reasons = {(x["path"], x["reason"]) for x in d}
    assert ("$.a", "value") in reasons and ("$.b", "length") in reasons
    assert ("$.c", "missing_right") not in reasons and ("$.d", "missing_left") in reasons
    assert ("$.t", "type_or_value") in reasons
    assert m.diff({"k": 1}, {}) == [{"path": "$.k", "left": 1, "right": None, "reason": "missing_right"}]
    assert m.diff([1, {"z": 1}], [1, {"z": 1}]) == []


def _manifest(tmp_path, name, body):
    p = tmp_path / name
    p.write_text(json.dumps(body), encoding="utf-8-sig")
    return p


def test_compare_main_match_and_difference(tmp_path, capsys):
    left = _manifest(tmp_path, "l.json", {"kind": "html", "generated_utc": "t1", "h": ["A"]})
    right = _manifest(tmp_path, "r.json", {"kind": "html", "generated_utc": "t2", "h": ["A"]})
    out = tmp_path / "o" / "rec.json"
    with pytest.raises(SystemExit) as exc:
        run_main(CMP, [str(left), str(right), "--output", str(out)])
    assert exc.value.code == 0
    rec = json.loads(out.read_text(encoding="utf-8"))
    assert rec["result"] == "PASS" and rec["semantic_match"] is True
    other = _manifest(tmp_path, "x.json", {"kind": "html", "h": ["B"]})
    with pytest.raises(SystemExit) as exc:
        run_main(CMP, [str(left), str(other), "--output", str(out)])
    assert exc.value.code == 1
    with pytest.raises(SystemExit) as exc:
        run_main(CMP, [str(left), str(other), "--output", str(out), "--allow-difference"])
    assert exc.value.code == 0
    assert json.loads(out.read_text(encoding="utf-8"))["difference_count"] == 1
    capsys.readouterr()
