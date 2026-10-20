"""House-style post-processing and QA gate (tools/house_style)."""
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
pytest.importorskip("pptx")

from tools.house_style import qps_post, qps_qa

ROOT = Path(__file__).resolve().parents[1]
STARTER = ROOT / "tools" / "house_style" / "templates" / "QPS_House_starter.docx"


class Meta:
    title = "VCR Handover Summary"
    subtitle = "Handover report"
    project = "MINERVA - Cryoplant - QPS - VCR"
    author = "Test Author"
    ref = "SCK CEN/12345678"
    third_party_ref = ""
    version = "0.1"
    status = "Draft"
    classification = "Restricted"
    date = "2026-10-11"
    short = False


def _house_doc(tmp_path, heading="Requirements"):
    d = docx.Document(str(STARTER))
    d.add_paragraph(heading, style="Heading 1")
    d.add_paragraph("Body text.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "ID", "Item"
    src = tmp_path / "src.docx"
    d.save(str(src))
    return src


def test_bare_starter_document_is_not_publishable(tmp_path):
    rep = qps_qa.check_docx(str(_house_doc(tmp_path)))
    rules = {i["rule"] for i in rep.items if i["level"] == "FAIL"}
    assert not rep.ok
    assert {"front matter", "title page", "header/footer"} <= rules


def test_post_adds_minerva_front_matter_and_footer_fields(tmp_path):
    out = tmp_path / "out.docx"
    qps_post.post_docx(str(_house_doc(tmp_path)), str(out), Meta, keep=True)
    rep = qps_qa.check_docx(str(out))
    assert rep.ok, rep.markdown()
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    for heading in qps_qa.FRONT_MATTER:
        assert heading in texts
    assert "SCK CEN Ref.: SCK CEN/12345678" in texts


def test_fix_strips_typed_heading_numbers_and_off_palette_colour(tmp_path):
    d = docx.Document(str(STARTER))
    d.add_paragraph("2. Quality Controls", style="Heading 1")
    run = d.add_paragraph().add_run("red text")
    run.font.color.rgb = docx.shared.RGBColor(0xFF, 0x00, 0x00)
    src = tmp_path / "src.docx"
    d.save(str(src))
    out = tmp_path / "fixed.docx"
    qps_post.post_docx(str(src), str(tmp_path / "staged.docx"), Meta, keep=True)
    qps_qa.fix_docx(str(tmp_path / "staged.docx"), str(out), Meta)
    rep = qps_qa.check_docx(str(out))
    assert rep.ok, rep.markdown()
    fixed = docx.Document(str(out))
    assert "Quality Controls" in [p.text for p in fixed.paragraphs]


@pytest.mark.parametrize("value,ok", [("562873", True), ("984A9C", True), ("EEE4F5", True),
                                      ("333333", True), ("FF0000", False), ("1D7A5F", False)])
def test_palette_rules(value, ok):
    assert qps_qa.colour_ok(value) is ok


def test_nearest_palette_colour():
    assert qps_qa.nearest("003366") == "034694"


@pytest.mark.skipif(not Path(qps_post.PPTX_REF).exists(), reason="corporate slide master is built locally")
def test_foreign_deck_is_rehosted_and_publishable(tmp_path):
    from pptx import Presentation
    src = tmp_path / "foreign.pptx"
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Scope"
    s.placeholders[1].text = "One\nTwo"
    prs.save(str(src))
    out = tmp_path / "out.pptx"
    qps_qa.fix_pptx(str(src), str(out), Meta)
    assert qps_qa.check_pptx(str(out)).ok


# ------------------------------------------------- slide kit regressions (K1-K6) ----
# Synthetic decks only: the corporate slide master is never committed.

ECM_ID = "87654321"  # synthetic
CUSTOM_PROPS = [
    ("ID", "lpwstr", ECM_ID), ("Name", "lpwstr", "Source Deck Title"),
    ("HyperLink", "lpwstr", f"https://ecm.example/open/{ECM_ID}"),
    ("Common Attributes_Reference Number", "lpwstr", f"SCK CEN/{ECM_ID}"),
    ("Common Attributes_Short Reference", "lpwstr", f"SCK CEN/{ECM_ID}"),
    ("Common Attributes_Author_Author Name", "lpwstr", "Someone Else"),
    ("Common Attributes_Information Security Classification", "lpwstr", "Restricted"),
    ("Infrastructure Attributes_Infrastructure_System", "lpwstr", "kept"),
    ("CreateDate", "filetime", "2021-11-30T19:34:52Z"),
]
CUSTOM_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    + "".join(f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i}" name="{k}">'
              f"<vt:{t}>{v}</vt:{t}></property>" for i, (k, t, v) in enumerate(CUSTOM_PROPS, start=2))
    + "</Properties>")
STALE_APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    "<TotalTime>28898</TotalTime><Application>Microsoft Office PowerPoint</Application>"
    "<PresentationFormat>Widescreen</PresentationFormat><Slides>38</Slides><Notes>23</Notes>"
    '<TitlesOfParts><vt:vector size="1" baseType="lpstr"><vt:lpstr>Source Slide 1</vt:lpstr></vt:vector>'
    "</TitlesOfParts><Company>SCK CEN</Company></Properties>")


def _textbox(shapes, text):
    """Text box on a master or layout (python-pptx only adds shapes to slides)."""
    from pptx.oxml.shapes.autoshape import CT_Shape
    sid = shapes._next_shape_id
    shapes._spTree.append(CT_Shape.new_textbox_sp(sid, f"TextBox {sid}", 100000, 6400000, 2000000, 300000))
    tb = shapes[len(shapes) - 1]
    tb.text_frame.text = text
    return tb


def _add_master(prs):
    """A second slide master (copy of the first, no layouts), as in the four-master corporate deck."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.packuri import PackURI
    from pptx.oxml.ns import qn
    from pptx.parts.slide import SlideMasterPart
    first = prs.slide_masters[0].part
    part = SlideMasterPart.load(PackURI("/ppt/slideMasters/slideMaster2.xml"), first.content_type,
                                prs.part.package, etree.tostring(first._element))
    part.relate_to(first.part_related_by(RT.THEME), RT.THEME)
    for lid in list(part._element.find(qn("p:sldLayoutIdLst"))):
        lid.getparent().remove(lid)
    rid = prs.part.relate_to(part, RT.SLIDE_MASTER)
    etree.SubElement(prs.part._element.find(qn("p:sldMasterIdLst")), qn("p:sldMasterId"),
                     {"id": "2147483700", qn("r:id"): rid})


def _ecm_deck(path, masters=2, slides=3):
    """A deck cut from an ECM document: identity in custom.xml, stale app.xml, furniture on every master."""
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    prs = Presentation()
    prs.slide_layouts[0].name, prs.slide_layouts[1].name = "Titel Slide 1", "Text Slide"
    while len(prs.slide_masters) < masters:
        _add_master(prs)
    for m in prs.slide_masters:
        _textbox(m.shapes, f"SCK CEN/{ECM_ID}")
        _textbox(m.shapes, "ISC: Restricted")
    _textbox(prs.slide_layouts[1].shapes, "ISC: Restricted")
    for n in range(slides):
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = f"Slide {n + 1}"
    pkg = prs.part.package
    pkg.relate_to(Part(PackURI("/docProps/custom.xml"),
                       "application/vnd.openxmlformats-officedocument.custom-properties+xml",
                       pkg, CUSTOM_XML.encode()), RT.CUSTOM_PROPERTIES)
    qps_post._docprops(prs, "app.xml")._blob = STALE_APP_XML.encode()
    prs.core_properties.author = "Someone Else"
    prs.save(str(path))
    return path


def _zip_text(path, name):
    import zipfile
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf-8")


def test_finish_deck_drops_ecm_identity_and_stale_app_props(tmp_path):  # K1, K2
    from pptx import Presentation
    prs = Presentation(str(_ecm_deck(tmp_path / "ecm.pptx")))
    qps_post.finish_deck(prs, Meta)
    out = tmp_path / "out.pptx"
    prs.save(str(out))
    custom = _zip_text(out, "docProps/custom.xml")
    assert ECM_ID not in custom and "Someone Else" not in custom and "CreateDate" not in custom
    props = qps_qa.custom_props(Presentation(str(out)))
    assert props["ID"] == props["HyperLink"] == ""
    assert props["Name"] == Meta.title
    assert props["Common Attributes_Short Reference"] == props["Common Attributes_Reference Number"] == Meta.ref
    assert props["Common Attributes_Author_Author Name"] == Meta.author
    assert props["Infrastructure Attributes_Infrastructure_System"] == "kept"
    app = _zip_text(out, "docProps/app.xml")
    assert "<Slides>4</Slides>" in app  # three slides plus the cover
    assert "TitlesOfParts" not in app and "TotalTime" not in app and "<Company>SCK CEN</Company>" in app
    assert Presentation(str(out)).core_properties.author == Meta.author


def test_finish_deck_keeps_the_ecm_link_when_ref_names_that_document(tmp_path):  # K1
    from pptx import Presentation
    prs = Presentation(str(_ecm_deck(tmp_path / "ecm.pptx")))

    class SameDoc(Meta):
        ref = f"SCK CEN/{ECM_ID}"
    qps_post.finish_deck(prs, SameDoc)
    props = qps_qa.custom_props(prs)
    assert props["ID"] == ECM_ID and ECM_ID in props["HyperLink"] and "CreateDate" in props
    assert props["Name"] == Meta.title and props["Common Attributes_Author_Author Name"] == Meta.author


def test_finish_deck_fills_furniture_on_every_master(tmp_path):  # K3
    from pptx import Presentation
    prs = Presentation(str(_ecm_deck(tmp_path / "ecm.pptx")))
    qps_post.finish_deck(prs, Meta)
    texts = {(str(part), role): shp.text_frame.text for part, role, shp in qps_post.furniture(prs)}
    assert {p for p, _ in texts} >= {"/ppt/slideMasters/slideMaster1.xml", "/ppt/slideMasters/slideMaster2.xml",
                                     "/ppt/slideLayouts/slideLayout2.xml"}
    assert {t for (_, role), t in texts.items() if role == "ref"} == {Meta.ref}
    assert {t for (_, role), t in texts.items() if role == "isc"} == {f"ISC: {Meta.classification}"}


def test_finish_deck_refills_a_reference_another_tool_wrote_on_every_master(tmp_path):  # K3
    from pptx import Presentation
    prs = Presentation(str(_ecm_deck(tmp_path / "ecm.pptx")))
    refs = [shp for _, role, shp in qps_post.furniture(prs) if role == "ref"]
    for shp in refs:
        qps_post._set_text(shp, "ICD-EL-001 rev 1.2")  # not SCK CEN/<digits> any more
    refs[0].name = "QPS Ref"  # only the first master was named by the old build_templates.py
    qps_post.finish_deck(prs, Meta)
    assert [shp.text_frame.text for _, role, shp in qps_post.furniture(prs) if role == "ref"] == [Meta.ref] * 2


def test_finish_deck_without_ref_does_not_keep_the_source_reference(tmp_path):  # K1, K3
    from pptx import Presentation
    prs = Presentation(str(_ecm_deck(tmp_path / "ecm.pptx")))

    class NoRef(Meta):
        ref = ""
    qps_post.finish_deck(prs, NoRef)
    assert {shp.text_frame.text for _, role, shp in qps_post.furniture(prs) if role == "ref"} == {"SCK CEN/"}


def test_qa_flags_foreign_ecm_identity_and_mixed_references(tmp_path):  # K1, K3 gate
    from pptx import Presentation
    prs = Presentation(str(_ecm_deck(tmp_path / "ecm.pptx")))
    rep = qps_qa.Report("ecm.pptx")
    qps_qa._check_identity(prs, rep)
    assert not rep.items  # the ECM's own deck: ID and footer agree
    _, _, first_ref = next(f for f in qps_post.furniture(prs) if f[1] == "ref")
    qps_post._set_text(first_ref, Meta.ref)  # old kit: master 1 refilled, master 2 and custom.xml left alone
    rep = qps_qa.Report("old-kit.pptx")
    qps_qa._check_identity(prs, rep)
    assert {i["rule"] for i in rep.items if i["level"] == "FAIL"} == {"identity", "reference"}
    qps_post.finish_deck(prs, Meta)
    rep = qps_qa.Report("fixed.pptx")
    qps_qa._check_identity(prs, rep)
    assert not rep.items


def test_build_slides_strips_identity_from_the_template(tmp_path):  # K1, K2, K7
    from pptx import Presentation

    from tools.house_style import build_templates
    deck = _ecm_deck(tmp_path / "source.pptx")
    sha = build_templates.build_slides(str(deck), str(tmp_path))
    built = tmp_path / "QPS_House.pptx"
    assert ECM_ID not in _zip_text(built, "docProps/custom.xml")
    assert "<Slides>0</Slides>" in _zip_text(built, "docProps/app.xml")
    for part in ("ppt/slideMasters/slideMaster1.xml", "ppt/slideMasters/slideMaster2.xml"):
        assert ECM_ID not in _zip_text(built, part)
    prs = Presentation(str(built))
    assert len(prs.slides) == 0 and prs.core_properties.author == ""
    named = {shp.name for m in prs.slide_masters for shp in m.shapes if qps_post.furniture_role(shp)}
    assert named >= {"QPS Ref", "QPS ISC"}
    assert f"sha256: {sha}" in (tmp_path / "QPS_House.source.txt").read_text(encoding="utf-8")


def test_layout_lookup_is_case_insensitive_with_fallbacks():  # K4
    from pptx import Presentation
    prs = Presentation()
    assert qps_post._layout(prs, "BLANK").name == "Blank"
    assert qps_post._layout(prs, "title and content").name == "Title and Content"
    next(lay for lay in prs.slide_layouts if lay.name == "Blank").name = "Blank 1"
    assert qps_post._layout(prs, qps_post.BLANK).name == "Blank 1"
    with pytest.raises(KeyError, match="Title Only"):
        qps_post._layout(prs, "No Such Layout")
    assert all(qps_post.is_blank_layout(n) for n in ("BLANK", "Blank", "Blank 2", "1_Blank 3", "Blank footer"))
    assert not qps_post.is_blank_layout("Text Slide")


@pytest.mark.parametrize("text,markers,found", [
    ("W363 burndown `status: TODO`", True, None),
    ("status: TODO", True, None),
    ('Cable cross-section "TBD" mm2', True, None),
    ("Cross-section TBD mm2", False, None),  # text box or table on a slide
    ("TODO: add the N-1 chart", True, "TODO"),
    ("Owner TBD", True, "TBD"),
    ("Lorem ipsum dolor", False, "Lorem ipsum"),
    ("SCK CEN Ref.: [ref]", False, "[ref]"),
    ("Mastodon todos", True, None),
])
def test_placeholder_text_spares_quoted_engineering_status(text, markers, found):  # K5
    assert qps_qa.placeholder_text(text, author_markers=markers) == found


def test_pdf_engine_prefers_soffice_then_office_on_windows(monkeypatch, tmp_path):  # K6
    monkeypatch.setattr(qps_post.sys, "platform", "win32")
    monkeypatch.setattr(qps_post.shutil, "which", lambda name: None)
    assert qps_post.pdf_engine() == "office"
    monkeypatch.setattr(qps_post.shutil, "which", lambda name: "C:/LibreOffice/program/soffice.exe")
    assert qps_post.pdf_engine() == "soffice"
    assert qps_post.pdf_engine("office") == "office"

    calls = []

    def fake_run(args, **kwargs):
        calls.append(args)
        Path(args[args.index("-Out") + 1]).write_bytes(b"%PDF-1.7")
        return type("Done", (), {"returncode": 0, "stdout": "", "stderr": ""})()
    monkeypatch.setattr(qps_post.subprocess, "run", fake_run)
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(b"")
    assert qps_post.export_pdf(str(deck), "office") == str(tmp_path / "deck.pdf")
    assert calls[0][0] == "powershell.exe" and str(qps_post.OFFICE_PDF) in calls[0]
