from __future__ import annotations

import hashlib
import json
import os
import shutil
import sys
from html.parser import HTMLParser
from importlib.metadata import version as package_version
from pathlib import Path
from typing import Any

from pptx import Presentation
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
A7 = ROOT / "abacus_render_pipeline" / "A7_executable_semantic_runtime"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(A7))

from codex.contract_governance.builder import build_artifacts, workbook_payload
from codex.contract_governance.io import content_hash, load_ssot
from semantic_delta_replay import run as run_semantic_replay

SSOT = ROOT / "contract_governance" / "ssot" / "abacus_contract_governance.yaml"
TUPLE_LEDGER = A7 / "data" / "semantic_tuple_ledger.json"
OUTPUT_DIR = HERE / "outputs"
RECEIPT_DIR = HERE / "receipts"
TIER = "internal"
REQUIRED_FORMATS = ("html", "pptx", "pdf", "markdown", "github_pages")


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)

    def text(self) -> str:
        return "\n".join(self.parts)


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _html_text(path: Path) -> str:
    parser = _TextExtractor()
    parser.feed(path.read_text(encoding="utf-8"))
    return parser.text()


def _pptx_text(path: Path) -> tuple[str, dict[str, Any]]:
    prs = Presentation(path)
    fragments: list[str] = []
    table_count = 0
    text_shape_count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_shape_count += 1
                fragments.append(shape.text_frame.text)
            if shape.has_table:
                table_count += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        fragments.append(cell.text)
    return "\n".join(fragments), {
        "postbuild_slide_count": len(prs.slides),
        "postbuild_table_count": table_count,
        "postbuild_text_shape_count": text_shape_count,
    }


def _pdf_text(path: Path) -> tuple[str, dict[str, Any]]:
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    sizes = [
        {
            "width_points": float(page.mediabox.width),
            "height_points": float(page.mediabox.height),
        }
        for page in reader.pages
    ]
    return "\n".join(pages), {
        "postbuild_page_count": len(reader.pages),
        "empty_page_count": sum(1 for text in pages if not text.strip()),
        "page_sizes": sizes,
    }


def _escape_md(value: object) -> str:
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def _render_markdown(payload: dict[str, Any], digest: str, path: Path) -> None:
    lines = [
        f"# {payload['package_id']} {str(payload['tier']).upper()} Governance Snapshot",
        "",
        f"Content hash (sha256): `{digest}`",
        "",
    ]
    for sheet in payload["sheets"]:
        lines.extend([f"## {sheet['name']}", ""])
        columns = [str(value) for value in sheet["columns"]]
        lines.append("| " + " | ".join(_escape_md(value) for value in columns) + " |")
        lines.append("| " + " | ".join("---" for _ in columns) + " |")
        for row in sheet["rows"]:
            lines.append(
                "| "
                + " | ".join(_escape_md(row.get(column, "")) for column in columns)
                + " |"
            )
        lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _semantic_tokens(payload: dict[str, Any], digest: str) -> list[str]:
    tokens = [str(payload["package_id"]), digest]
    for sheet in payload["sheets"]:
        tokens.append(str(sheet["name"]))
        for row in sheet["rows"]:
            req_id = row.get("Requirement ID")
            if req_id:
                tokens.append(str(req_id))
    return list(dict.fromkeys(tokens))


def _parity(text: str, expected_tokens: list[str]) -> dict[str, Any]:
    missing = [token for token in expected_tokens if token not in text]
    present = len(expected_tokens) - len(missing)
    return {
        "pass": not missing,
        "expected_token_count": len(expected_tokens),
        "present_token_count": present,
        "coverage": round(present / max(len(expected_tokens), 1), 6),
        "missing_tokens": missing,
    }


def _format_receipt(
    *,
    name: str,
    path: Path,
    renderer_version: str,
    text: str,
    telemetry: dict[str, Any],
    parity: dict[str, Any],
) -> dict[str, Any]:
    layout_pass = bool(telemetry.get("layout_pass", True))
    overflow_pass = bool(telemetry.get("overflow_pass", True))
    decision = "accept" if parity["pass"] and layout_pass and overflow_pass else "reject"
    return {
        "format": name,
        "artifact_relpath": path.relative_to(HERE).as_posix(),
        "artifact_sha256": sha256_path(path),
        "artifact_bytes": path.stat().st_size,
        "renderer_version": renderer_version,
        "telemetry": telemetry,
        "semantic_parity": parity,
        "extracted_text_chars": len(text),
        "decision": decision,
    }


def execute() -> dict[str, Any]:
    ssot = load_ssot(SSOT)
    payload = workbook_payload(ssot, TIER)
    canonical_digest = content_hash(payload)
    expected_tokens = _semantic_tokens(payload, canonical_digest)

    production_root = OUTPUT_DIR / "production"
    production = build_artifacts(ssot, production_root, TIER)
    native = production["renderer_telemetry"]
    if not isinstance(native, dict):
        raise ValueError("production renderer telemetry missing")

    html_path = Path(str(production["html"]))
    pptx_path = Path(str(production["pptx"]))
    pdf_path = Path(str(production["pdf"]))

    markdown_path = OUTPUT_DIR / f"{ssot.package_id}_{TIER}.md"
    _render_markdown(payload, canonical_digest, markdown_path)

    pages_path = OUTPUT_DIR / "github_pages" / "index.html"
    pages_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(html_path, pages_path)

    html_text = _html_text(html_path)
    pptx_text, pptx_postbuild = _pptx_text(pptx_path)
    pdf_text, pdf_postbuild = _pdf_text(pdf_path)
    markdown_text = markdown_path.read_text(encoding="utf-8")
    pages_text = _html_text(pages_path)

    pptx_native = dict(native.get("pptx", {}))
    pptx_native.update(pptx_postbuild)
    pdf_native = dict(native.get("pdf", {}))
    pdf_native.update(pdf_postbuild)
    pdf_native["overflow_pass"] = bool(pdf_native.get("overflow_pass")) and pdf_postbuild[
        "empty_page_count"
    ] == 0

    format_receipts = {
        "html": _format_receipt(
            name="html",
            path=html_path,
            renderer_version=f"jinja2/{package_version('jinja2')}+contract-governance",
            text=html_text,
            telemetry={
                "layout_pass": True,
                "overflow_pass": True,
                "method": "production_html_render_plus_semantic_coverage",
            },
            parity=_parity(html_text, expected_tokens),
        ),
        "pptx": _format_receipt(
            name="pptx",
            path=pptx_path,
            renderer_version=f"python-pptx/{package_version('python-pptx')}+contract-governance",
            text=pptx_text,
            telemetry=pptx_native,
            parity=_parity(pptx_text, expected_tokens),
        ),
        "pdf": _format_receipt(
            name="pdf",
            path=pdf_path,
            renderer_version=f"reportlab/{package_version('reportlab')}+contract-governance",
            text=pdf_text,
            telemetry=pdf_native,
            parity=_parity(pdf_text, expected_tokens),
        ),
        "markdown": _format_receipt(
            name="markdown",
            path=markdown_path,
            renderer_version="abacus-markdown/1.0.0",
            text=markdown_text,
            telemetry={
                "layout_pass": True,
                "overflow_pass": True,
                "method": "deterministic_table_serialization",
            },
            parity=_parity(markdown_text, expected_tokens),
        ),
        "github_pages": _format_receipt(
            name="github_pages",
            path=pages_path,
            renderer_version="abacus-github-pages-adapter/1.0.0",
            text=pages_text,
            telemetry={
                "layout_pass": True,
                "overflow_pass": True,
                "method": "byte-identical_production_html_deployment_surface",
                "source_html_sha256": sha256_path(html_path),
            },
            parity=_parity(pages_text, expected_tokens),
        ),
    }

    replay = run_semantic_replay()
    replay_pass = replay.get("status") == "PASS"
    formats_accept = all(
        format_receipts[name]["decision"] == "accept" for name in REQUIRED_FORMATS
    )
    parity_pass = all(format_receipts[name]["semantic_parity"]["pass"] for name in REQUIRED_FORMATS)
    pages_match_html = (
        format_receipts["github_pages"]["artifact_sha256"]
        == format_receipts["html"]["artifact_sha256"]
    )
    cross_format_parity = {
        "pass": parity_pass and pages_match_html,
        "formats_checked": list(REQUIRED_FORMATS),
        "canonical_content_sha256": canonical_digest,
        "github_pages_matches_html_bytes": pages_match_html,
        "semantic_token_count": len(expected_tokens),
    }

    decision = (
        "accept"
        if formats_accept and cross_format_parity["pass"] and replay_pass
        else "reject"
    )
    source_commit = os.environ.get("ABACUS_SOURCE_SHA") or os.environ.get("GITHUB_SHA") or "local"
    receipt = {
        "receipt_version": "A9.0",
        "publication_id": ssot.package_id,
        "source_commit": source_commit,
        "ssot_relpath": SSOT.relative_to(ROOT).as_posix(),
        "ssot_sha256": sha256_path(SSOT),
        "canonical_content_sha256": canonical_digest,
        "tuple_ledger_sha256": sha256_path(TUPLE_LEDGER),
        "theme": "contract-governance-default",
        "render_mode": TIER,
        "formats": format_receipts,
        "cross_format_parity": cross_format_parity,
        "semantic_replay": {
            "pass": replay_pass,
            "status": replay.get("status", "UNKNOWN"),
            "visited_count": replay.get("final_state", {}).get("visited_count"),
            "last_tuple": replay.get("final_state", {}).get("last_tuple"),
        },
        "decision": decision,
    }

    RECEIPT_DIR.mkdir(parents=True, exist_ok=True)
    (RECEIPT_DIR / "multiformat_execution_receipt.json").write_text(
        json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return receipt


if __name__ == "__main__":
    result = execute()
    print(json.dumps(result, indent=2, sort_keys=True))
    raise SystemExit(0 if result["decision"] == "accept" else 1)
