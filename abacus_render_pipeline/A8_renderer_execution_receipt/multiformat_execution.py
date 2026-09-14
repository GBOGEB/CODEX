from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "abacus_render_pipeline" / "A7_executable_semantic_runtime"))

from renderers.theme_runtime import SemanticThemeRuntime
from semantic_delta_replay import run as run_semantic_replay

HERE = Path(__file__).resolve().parent
DEFAULT_SSOT = HERE / "ssot" / "reference_publication.yaml"
TUPLE_LEDGER = ROOT / "abacus_render_pipeline" / "A7_executable_semantic_runtime" / "data" / "semantic_tuple_ledger.json"
OUTPUT_DIR = HERE / "outputs"
RECEIPT_DIR = HERE / "receipts"

PPTX_RENDERER_VERSION = "abacus-native-pptx/1.0.0"
PDF_RENDERER_VERSION = "abacus-native-pdf-reportlab/1.0.0"
PPTX_OUTPUT = OUTPUT_DIR / "production_render.pptx"
PDF_OUTPUT = OUTPUT_DIR / "production_render.pdf"
MULTIFORMAT_RECEIPT = RECEIPT_DIR / "multiformat_execution_receipt.json"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
PAGE_W_PT = SLIDE_W_IN * 72.0
PAGE_H_PT = SLIDE_H_IN * 72.0
FIXED_TIME = dt.datetime(2000, 1, 1, 0, 0, 0)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normalize_line(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    return value.lstrip("•- ").strip()


def _publication_lines(ssot: dict[str, Any]) -> list[str]:
    card = ssot["publication"]["semantic_card"]
    return [_normalize_line(str(card["title"]))] + [
        _normalize_line(str(line)) for line in card.get("body", [])
    ]


def _content_fingerprint(lines: list[str]) -> str:
    return sha256_bytes(("\n".join(lines) + "\n").encode("utf-8"))


def _rgb(value: str) -> RGBColor:
    raw = value.lstrip("#")
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def _canonicalize_zip(data: bytes) -> bytes:
    """Rewrite an OpenXML package with stable member ordering/timestamps."""
    source = io.BytesIO(data)
    target = io.BytesIO()
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as dst:
        for name in sorted(src.namelist()):
            original = src.getinfo(name)
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            dst.writestr(info, src.read(name))
    return target.getvalue()


def _extract_pptx_lines(data: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = _normalize_line(paragraph.text)
                if text:
                    lines.append(text)
    return lines


def _extract_pdf_lines(data: bytes) -> list[str]:
    reader = PdfReader(io.BytesIO(data))
    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        lines.extend(
            normalized
            for raw in text.splitlines()
            if (normalized := _normalize_line(raw))
        )
    return lines


def _pptx_native_telemetry(data: bytes, expected_lines: list[str]) -> tuple[dict[str, Any], dict[str, Any]]:
    prs = Presentation(io.BytesIO(data))
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    out_of_bounds: list[str] = []
    content_boxes: list[tuple[int, int, str]] = []
    height_margins_pt: list[float] = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            left = int(shape.left)
            top = int(shape.top)
            right = left + int(shape.width)
            bottom = top + int(shape.height)
            if left < 0 or top < 0 or right > slide_w or bottom > slide_h:
                out_of_bounds.append(f"s{slide_index}:shape{shape_index}")
            if not getattr(shape, "has_text_frame", False):
                continue
            paragraphs = [
                _normalize_line(paragraph.text)
                for paragraph in shape.text_frame.paragraphs
                if _normalize_line(paragraph.text)
            ]
            if not paragraphs:
                continue
            content_boxes.append((top, bottom, " | ".join(paragraphs)))
            sizes: list[float] = []
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        sizes.append(float(run.font.size.pt))
            max_size = max(sizes) if sizes else 18.0
            box_height_pt = float(shape.height) / 12700.0
            height_margins_pt.append(round(box_height_pt - max_size * 1.35, 3))

    content_boxes.sort(key=lambda row: row[0])
    overlaps = [
        f"{content_boxes[index - 1][2]} -> {content_boxes[index][2]}"
        for index in range(1, len(content_boxes))
        if content_boxes[index - 1][1] > content_boxes[index][0]
    ]
    extracted = _extract_pptx_lines(data)
    content_match = extracted == expected_lines
    min_height_margin = min(height_margins_pt) if height_margins_pt else -1.0

    telemetry = {
        "telemetry_class": "NATIVE_OPENXML_STRUCTURE_AND_FONT_GEOMETRY",
        "slide_count": len(prs.slides),
        "slide_width_emu": slide_w,
        "slide_height_emu": slide_h,
        "out_of_bounds_shapes": out_of_bounds,
        "content_box_overlap_count": len(overlaps),
        "content_box_overlaps": overlaps,
        "minimum_text_box_height_margin_pt": min_height_margin,
        "extracted_content_lines": extracted,
        "host_layout_engine_limitation": "PowerPoint host text reflow is not executed; telemetry is OpenXML structure plus renderer font geometry.",
    }
    checks = {
        "layout": {"pass": len(prs.slides) == 1 and not out_of_bounds and not overlaps},
        "overflow": {"pass": min_height_margin >= 0.0},
        "content": {"pass": content_match},
    }
    return telemetry, checks


def render_pptx(ssot: dict[str, Any]) -> tuple[bytes, dict[str, Any], dict[str, Any]]:
    publication = ssot["publication"]
    card = publication["semantic_card"]
    theme = SemanticThemeRuntime().resolve(card["type"], publication["theme"])
    expected_lines = _publication_lines(ssot)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = _rgb("#181421")
    background.line.fill.background()

    card_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.9),
        Inches(0.8),
        Inches(11.5),
        Inches(5.9),
    )
    card_shape.fill.solid()
    card_shape.fill.fore_color.rgb = _rgb(theme.background)
    card_shape.line.color.rgb = _rgb(theme.border)

    def add_text(
        x: float,
        y: float,
        w: float,
        h: float,
        text: str,
        size: float,
        bold: bool = False,
    ) -> None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(theme.text)

    add_text(1.3, 1.3, 10.7, 0.8, expected_lines[0], 26, True)
    for index, line in enumerate(expected_lines[1:]):
        add_text(1.45, 2.45 + index * 0.95, 10.35, 0.68, line, 20, False)

    prs.core_properties.author = "ABACUS governed renderer"
    prs.core_properties.last_modified_by = "ABACUS governed renderer"
    prs.core_properties.created = FIXED_TIME
    prs.core_properties.modified = FIXED_TIME
    prs.core_properties.title = str(publication["title"])

    raw = io.BytesIO()
    prs.save(raw)
    artifact = _canonicalize_zip(raw.getvalue())
    telemetry, checks = _pptx_native_telemetry(artifact, expected_lines)
    return artifact, telemetry, checks


def _pdf_native_telemetry(data: bytes, expected_lines: list[str], width_metrics: list[dict[str, float]]) -> tuple[dict[str, Any], dict[str, Any]]:
    reader = PdfReader(io.BytesIO(data))
    extracted = _extract_pdf_lines(data)
    pages = len(reader.pages)
    first_page = reader.pages[0] if pages else None
    width = float(first_page.mediabox.width) if first_page else 0.0
    height = float(first_page.mediabox.height) if first_page else 0.0
    min_margin = min((row["available_width_pt"] - row["text_width_pt"] for row in width_metrics), default=-1.0)
    content_match = extracted == expected_lines
    telemetry = {
        "telemetry_class": "NATIVE_REPORTLAB_TEXT_METRICS_PLUS_PDF_PAGE_STRUCTURE",
        "page_count": pages,
        "page_width_pt": round(width, 3),
        "page_height_pt": round(height, 3),
        "minimum_text_width_margin_pt": round(min_margin, 3),
        "line_width_metrics": width_metrics,
        "extracted_content_lines": extracted,
    }
    checks = {
        "layout": {
            "pass": pages == 1 and abs(width - PAGE_W_PT) < 0.5 and abs(height - PAGE_H_PT) < 0.5
        },
        "overflow": {"pass": min_margin >= 0.0},
        "content": {"pass": content_match},
    }
    return telemetry, checks


def render_pdf(ssot: dict[str, Any]) -> tuple[bytes, dict[str, Any], dict[str, Any]]:
    publication = ssot["publication"]
    card = publication["semantic_card"]
    theme = SemanticThemeRuntime().resolve(card["type"], publication["theme"])
    expected_lines = _publication_lines(ssot)

    target = io.BytesIO()
    pdf = canvas.Canvas(
        target,
        pagesize=(PAGE_W_PT, PAGE_H_PT),
        invariant=1,
        pageCompression=1,
    )
    pdf.setAuthor("ABACUS governed renderer")
    pdf.setCreator(PDF_RENDERER_VERSION)
    pdf.setTitle(str(publication["title"]))

    pdf.setFillColor(HexColor("#181421"))
    pdf.rect(0, 0, PAGE_W_PT, PAGE_H_PT, stroke=0, fill=1)

    card_x = 0.9 * 72.0
    card_y = 0.8 * 72.0
    card_w = 11.5 * 72.0
    card_h = 5.9 * 72.0
    pdf.setFillColor(HexColor(theme.background))
    pdf.setStrokeColor(HexColor(theme.border))
    pdf.roundRect(card_x, card_y, card_w, card_h, 10, stroke=1, fill=1)

    width_metrics: list[dict[str, float]] = []

    def draw_line(x_in: float, y_from_top_in: float, text: str, size: float, font: str) -> None:
        x = x_in * 72.0
        y = PAGE_H_PT - y_from_top_in * 72.0
        available = (SLIDE_W_IN - x_in - 1.3) * 72.0
        text_width = stringWidth(text, font, size)
        width_metrics.append(
            {
                "font_size_pt": float(size),
                "text_width_pt": round(float(text_width), 3),
                "available_width_pt": round(float(available), 3),
            }
        )
        pdf.setFont(font, size)
        pdf.setFillColor(HexColor(theme.text))
        pdf.drawString(x, y, text)

    draw_line(1.3, 1.85, expected_lines[0], 26, "Helvetica-Bold")
    for index, line in enumerate(expected_lines[1:]):
        draw_line(1.45, 2.95 + index * 0.95, line, 20, "Helvetica")

    pdf.showPage()
    pdf.save()
    artifact = target.getvalue()
    telemetry, checks = _pdf_native_telemetry(artifact, expected_lines, width_metrics)
    return artifact, telemetry, checks


def execute_multiformat(ssot_path: Path = DEFAULT_SSOT) -> dict[str, Any]:
    ssot_bytes = ssot_path.read_bytes()
    ssot = yaml.safe_load(ssot_bytes)
    if not isinstance(ssot, dict) or "publication" not in ssot:
        raise ValueError("SSOT must contain a publication mapping")

    expected_lines = _publication_lines(ssot)
    expected_content_sha = _content_fingerprint(expected_lines)

    pptx_bytes, pptx_telemetry, pptx_checks = render_pptx(ssot)
    pdf_bytes, pdf_telemetry, pdf_checks = render_pdf(ssot)

    pptx_lines = _extract_pptx_lines(pptx_bytes)
    pdf_lines = _extract_pdf_lines(pdf_bytes)
    pptx_content_sha = _content_fingerprint(pptx_lines)
    pdf_content_sha = _content_fingerprint(pdf_lines)

    replay = run_semantic_replay()
    replay_pass = replay.get("status") == "PASS"

    pptx_decision = "accept" if all(item["pass"] for item in pptx_checks.values()) else "reject"
    pdf_decision = "accept" if all(item["pass"] for item in pdf_checks.values()) else "reject"
    parity_pass = (
        pptx_content_sha == expected_content_sha
        and pdf_content_sha == expected_content_sha
        and pptx_content_sha == pdf_content_sha
    )
    overall_accept = pptx_decision == "accept" and pdf_decision == "accept" and parity_pass and replay_pass

    publication = ssot["publication"]
    source_commit = os.environ.get("ABACUS_SOURCE_SHA") or os.environ.get("GITHUB_SHA") or "local"
    receipt = {
        "receipt_version": "A8.2",
        "publication_id": publication["id"],
        "source_commit": source_commit,
        "theme": publication["theme"],
        "render_mode": publication["render_mode"],
        "required_formats": ["pptx", "pdf"],
        "ssot_sha256": sha256_bytes(ssot_bytes),
        "tuple_ledger_sha256": sha256_path(TUPLE_LEDGER),
        "expected_content_sha256": expected_content_sha,
        "artifacts": {
            "pptx": {
                "artifact_sha256": sha256_bytes(pptx_bytes),
                "renderer_version": PPTX_RENDERER_VERSION,
                "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "native_telemetry": pptx_telemetry,
                "checks": pptx_checks,
                "decision": pptx_decision,
            },
            "pdf": {
                "artifact_sha256": sha256_bytes(pdf_bytes),
                "renderer_version": PDF_RENDERER_VERSION,
                "mime_type": "application/pdf",
                "native_telemetry": pdf_telemetry,
                "checks": pdf_checks,
                "decision": pdf_decision,
            },
        },
        "cross_format_parity": {
            "pass": parity_pass,
            "pptx_content_sha256": pptx_content_sha,
            "pdf_content_sha256": pdf_content_sha,
        },
        "semantic_replay": {
            "pass": replay_pass,
            "status": replay.get("status", "UNKNOWN"),
            "tuple_count": replay.get("final_state", {}).get("visited_count"),
        },
        "decision": "accept" if overall_accept else "reject",
    }

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RECEIPT_DIR.mkdir(parents=True, exist_ok=True)
    PPTX_OUTPUT.write_bytes(pptx_bytes)
    PDF_OUTPUT.write_bytes(pdf_bytes)
    MULTIFORMAT_RECEIPT.write_text(
        json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return receipt


if __name__ == "__main__":
    result = execute_multiformat()
    print(json.dumps(result, indent=2, sort_keys=True))
    raise SystemExit(0 if result["decision"] == "accept" else 1)
