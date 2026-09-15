from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

import yaml
from jsonschema import Draft202012Validator
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfbase.pdfmetrics import stringWidth

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SCHEMA = HERE / "multiformat_receipt_schema.json"
SSOT = HERE / "ssot" / "reference_publication.yaml"
TUPLE_LEDGER = ROOT / "abacus_render_pipeline" / "A7_executable_semantic_runtime" / "data" / "semantic_tuple_ledger.json"
PPTX_ARTIFACT = HERE / "outputs" / "production_render.pptx"
PDF_ARTIFACT = HERE / "outputs" / "production_render.pdf"

PAGE_W_PT = 13.333 * 72.0
PAGE_H_PT = 7.5 * 72.0


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _normalize_line(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    return value.lstrip("•- ").strip()


def _expected_lines() -> list[str]:
    doc = yaml.safe_load(SSOT.read_text(encoding="utf-8"))
    card = doc["publication"]["semantic_card"]
    return [_normalize_line(str(card["title"]))] + [
        _normalize_line(str(line)) for line in card.get("body", [])
    ]


def _fingerprint(lines: list[str]) -> str:
    return sha256_bytes(("\n".join(lines) + "\n").encode("utf-8"))


def _pptx_facts() -> dict[str, Any]:
    prs = Presentation(str(PPTX_ARTIFACT))
    lines: list[str] = []
    out_of_bounds = 0
    height_margins: list[float] = []
    content_boxes: list[tuple[int, int]] = []
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)

    for slide in prs.slides:
        for shape in slide.shapes:
            left = int(shape.left)
            top = int(shape.top)
            right = left + int(shape.width)
            bottom = top + int(shape.height)
            if left < 0 or top < 0 or right > slide_w or bottom > slide_h:
                out_of_bounds += 1
            if not getattr(shape, "has_text_frame", False):
                continue
            shape_lines = [
                _normalize_line(paragraph.text)
                for paragraph in shape.text_frame.paragraphs
                if _normalize_line(paragraph.text)
            ]
            if not shape_lines:
                continue
            lines.extend(shape_lines)
            content_boxes.append((top, bottom))
            sizes: list[float] = []
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        sizes.append(float(run.font.size.pt))
            max_size = max(sizes) if sizes else 18.0
            height_margins.append(float(shape.height) / 12700.0 - max_size * 1.35)

    content_boxes.sort(key=lambda item: item[0])
    overlaps = sum(
        1
        for index in range(1, len(content_boxes))
        if content_boxes[index - 1][1] > content_boxes[index][0]
    )
    return {
        "slide_count": len(prs.slides),
        "out_of_bounds_shapes": out_of_bounds,
        "content_box_overlap_count": overlaps,
        "minimum_text_box_height_margin_pt": min(height_margins) if height_margins else -1.0,
        "lines": lines,
    }


def _pdf_facts(expected: list[str]) -> dict[str, Any]:
    reader = PdfReader(str(PDF_ARTIFACT))
    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        lines.extend(
            normalized
            for raw in text.splitlines()
            if (normalized := _normalize_line(raw))
        )
    first_page = reader.pages[0] if reader.pages else None
    width = float(first_page.mediabox.width) if first_page else 0.0
    height = float(first_page.mediabox.height) if first_page else 0.0

    width_rows = [(expected[0], "Helvetica-Bold", 26.0)] + [
        (line, "Helvetica", 20.0) for line in expected[1:]
    ]
    available_widths = [
        (13.333 - 1.3 - 1.3) * 72.0,
        *[(13.333 - 1.45 - 1.3) * 72.0 for _ in expected[1:]],
    ]
    margins = [
        available - stringWidth(text, font, size)
        for (text, font, size), available in zip(width_rows, available_widths)
    ]
    return {
        "page_count": len(reader.pages),
        "page_width_pt": width,
        "page_height_pt": height,
        "minimum_text_width_margin_pt": min(margins) if margins else -1.0,
        "lines": lines,
    }


def load_and_validate_multiformat_receipt(receipt_path: Path) -> dict[str, Any]:
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    schema = json.loads(SCHEMA.read_text(encoding="utf-8"))
    errors = sorted(Draft202012Validator(schema).iter_errors(receipt), key=lambda e: list(e.path))
    if errors:
        raise ValueError("multiformat receipt schema validation failed: " + "; ".join(error.message for error in errors))

    expected = _expected_lines()
    expected_content_sha = _fingerprint(expected)
    pptx = _pptx_facts()
    pdf = _pdf_facts(expected)
    pptx_content_sha = _fingerprint(pptx["lines"])
    pdf_content_sha = _fingerprint(pdf["lines"])

    mismatches: list[str] = []
    expected_hashes = {
        "ssot_sha256": sha256_path(SSOT),
        "tuple_ledger_sha256": sha256_path(TUPLE_LEDGER),
        "expected_content_sha256": expected_content_sha,
    }
    for field, actual in expected_hashes.items():
        if receipt.get(field) != actual:
            mismatches.append(f"{field}: receipt={receipt.get(field)} actual={actual}")

    artifact_hashes = {
        "pptx": sha256_path(PPTX_ARTIFACT),
        "pdf": sha256_path(PDF_ARTIFACT),
    }
    for fmt, actual in artifact_hashes.items():
        observed = receipt["artifacts"][fmt]["artifact_sha256"]
        if observed != actual:
            mismatches.append(f"{fmt}.artifact_sha256: receipt={observed} actual={actual}")

    if mismatches:
        raise ValueError("multiformat content-address validation failed: " + "; ".join(mismatches))

    recomputed_checks = {
        "pptx": {
            "layout": pptx["slide_count"] == 1 and pptx["out_of_bounds_shapes"] == 0 and pptx["content_box_overlap_count"] == 0,
            "overflow": pptx["minimum_text_box_height_margin_pt"] >= 0.0,
            "content": pptx["lines"] == expected,
        },
        "pdf": {
            "layout": pdf["page_count"] == 1 and abs(pdf["page_width_pt"] - PAGE_W_PT) < 0.5 and abs(pdf["page_height_pt"] - PAGE_H_PT) < 0.5,
            "overflow": pdf["minimum_text_width_margin_pt"] >= 0.0,
            "content": pdf["lines"] == expected,
        },
    }
    for fmt, checks in recomputed_checks.items():
        for name, passed in checks.items():
            observed = bool(receipt["artifacts"][fmt]["checks"][name]["pass"])
            if observed != passed:
                raise ValueError(
                    f"multiformat native check mismatch for {fmt}.{name}: receipt={observed} actual={passed}"
                )
        expected_decision = "accept" if all(checks.values()) else "reject"
        if receipt["artifacts"][fmt]["decision"] != expected_decision:
            raise ValueError(
                f"multiformat artifact decision mismatch for {fmt}: receipt={receipt['artifacts'][fmt]['decision']} expected={expected_decision}"
            )

    parity_pass = (
        pptx_content_sha == expected_content_sha
        and pdf_content_sha == expected_content_sha
        and pptx_content_sha == pdf_content_sha
    )
    parity = receipt["cross_format_parity"]
    if bool(parity["pass"]) != parity_pass:
        raise ValueError(
            f"cross-format parity mismatch: receipt={parity['pass']} actual={parity_pass}"
        )
    if parity["pptx_content_sha256"] != pptx_content_sha or parity["pdf_content_sha256"] != pdf_content_sha:
        raise ValueError("cross-format content fingerprints do not match independently extracted artifacts")

    replay_pass = bool(receipt["semantic_replay"]["pass"])
    expected_decision = "accept" if (
        all(receipt["artifacts"][fmt]["decision"] == "accept" for fmt in ("pptx", "pdf"))
        and parity_pass
        and replay_pass
    ) else "reject"
    if receipt["decision"] != expected_decision:
        raise ValueError(
            f"multiformat receipt decision inconsistent: receipt={receipt['decision']} expected={expected_decision}"
        )

    return receipt


if __name__ == "__main__":
    path = HERE / "receipts" / "multiformat_execution_receipt.json"
    validated = load_and_validate_multiformat_receipt(path)
    print(json.dumps({"status": "PASS", "decision": validated["decision"], "formats": validated["required_formats"]}, indent=2))
